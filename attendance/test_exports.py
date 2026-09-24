import csv
import io
from datetime import date

from django.contrib.auth.models import User
from rest_framework.test import APITestCase

from attendance.exports import safe_cell, safe_rows
from attendance.models import AttendanceRecord, AttendanceSession
from buses.models import Bus
from students.models import Student

EVIL = "=1+1"


class SafeCellTests(APITestCase):
    """Audit item 1.8: exported text must never be readable as a spreadsheet formula."""

    def test_formula_starters_get_a_quote(self):
        for text in ["=1+1", "+91 98765", "-2+3", "@SUM(A1)", "\t=cmd", "\r=cmd"]:
            self.assertEqual(safe_cell(text), "'" + text, repr(text))

    def test_normal_values_are_unchanged(self):
        for value in ["Alice", "R950", "-", "", "85.0%", 12, 0, None]:
            self.assertEqual(safe_cell(value), value)

    def test_safe_rows_covers_every_cell(self):
        self.assertEqual(safe_rows([["a", "=x"], ["+y", 3]]), [["a", "'=x"], ["'+y", 3]])


class ExportInjectionTests(APITestCase):
    def setUp(self):
        self.bus = Bus.objects.create(
            bus_number="B77", route="Route 77", rfid_uid="RFID-77", departure_time="08:00",
            length_m="10.00", width_m="2.50",
        )
        self.student = Student.objects.create(
            name=EVIL, roll_number="=CMD|calc", phone="+919876543210", bus=self.bus,
        )
        session = AttendanceSession.objects.create(bus=self.bus, date=date.today(), slot="MORNING")
        AttendanceRecord.objects.create(
            session=session, person_type="STUDENT", student=self.student, status="PRESENT",
        )
        staff = User.objects.create_user("exp_staff", password="x")
        staff.profile.role = "STAFF"
        staff.profile.save()
        self.client.force_authenticate(staff)

    def assertNoFormulaCells(self, text):
        for row in csv.reader(io.StringIO(text.lstrip("\ufeff"))):
            for cell in row:
                self.assertFalse(cell.startswith(("=", "+", "@")), f"formula-like cell: {cell!r}")

    def test_history_csv_export_has_no_formula_cells(self):
        r = self.client.get(f"/api/attendance/export/?bus={self.bus.pk}&filetype=csv")
        self.assertEqual(r.status_code, 200)
        text = r.content.decode("utf-8")
        self.assertIn("'=1+1", text)
        self.assertNoFormulaCells(text)

    def test_history_xlsx_export_stores_names_as_text(self):
        from openpyxl import load_workbook

        r = self.client.get(f"/api/attendance/export/?bus={self.bus.pk}&filetype=xlsx")
        self.assertEqual(r.status_code, 200)
        ws = load_workbook(io.BytesIO(r.content)).active
        cells = [c for row in ws.iter_rows() for c in row if c.value == EVIL]
        self.assertEqual(len(cells), 1)
        self.assertEqual(cells[0].data_type, "s")  # text, not a formula

    def test_student_report_csv_has_no_formula_cells(self):
        r = self.client.get(f"/api/attendance/report/?bus={self.bus.pk}&filetype=csv")
        self.assertEqual(r.status_code, 200)
        text = r.content.decode("utf-8")
        self.assertIn("'=1+1", text)
        self.assertIn("'+919876543210", text)
        self.assertNoFormulaCells(text)
